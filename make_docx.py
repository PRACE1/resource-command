from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

def create_doc():
    doc = Document()
    
    # Page 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Republic of Botswana\n\n").bold = True
    p.add_run("Companies Act (CAP 42:01)\n\n").bold = True
    p.add_run("(As Amended)\n\n").bold = True
    p.add_run("A Private Company Having Share Capital\n\n").bold = True
    p.add_run("THE CONSTITUTION\n\nof\n\n").bold = True
    p.add_run("Company Name: Kgosi Sovereign Holdings Proprietary Limited\n\n").bold = True
    p.add_run("Company Registration Number: [Pending]\n\n\n\n").bold = True
    p.add_run("Adopted By Special Resolution Passed on: 09 / 06 / 2026").bold = True

    doc.add_page_break()

    # Page 2
    p2 = doc.add_paragraph()
    p2.add_run("A. CONSTITUTION OF A PRIVATE COMPANY LIMITED BY SHARES\n").bold = True
    p2.add_run("(sections 37(2), 40(b), 109, 129(1), and 156)\n\n").italic = True
    p2.add_run("COMPANY NAME: Kgosi Sovereign Holdings Proprietary Limited\n\n")

    sections = [
        ("1. Interpretation", "In this constitution –\n'Act' means the Companies Act;\n'beneficial owner' has the meaning and understanding ascribed to it under section 2 of the Financial Intelligence Act;\n'ultimate effective control' has the meaning and understanding ascribed to it under section 2 of the Financial Intelligence Act;\n'Founder' means Kennedy Jr. Zibo Thebe;\n'IP Assets' means all intellectual property owned or licensed by the company, including but not limited to zero-knowledge royalty verification systems, IoT transit tracking integrations, and related cryptographic protocols;\n'SAFE' means a Simple Agreement for Future Equity."),
        ("2. Issue of new shares mandate", "(1) New shares shall be issued in accordance with section 50 of the Act with the pre-emptive rights provided for in section 52.\n(2) At incorporation, the Founder holds 100,000 ordinary shares, being 100% of all issued shares. A further 35,000 shares are authorised but unissued and reserved for future allotment by the board to investors, strategic partners, and employees as and when the Founder determines. These reserved shares carry no rights until formally issued by board resolution.\n(3) SAFE Instruments: The board may issue SAFE instruments. SAFEs shall convert to shares upon a qualifying financing event. The aggregate principal amount of all outstanding SAFEs shall not exceed USD 500,000 without a Special Resolution."),
        ("3. Transfer of shares", "(1) Freedom to Transfer is Qualified: Every change in the ownership of shares in the capital of the company shall be subject to the following limitations and restrictions.\n(2) Pre-emptive Provisions: No share in the capital of the company shall be sold or transferred by any shareholder unless and until the rights of pre-emption hereinafter conferred have been exhausted.\n(3) Transfer Notice and Fair Price: It shall be required that –\n(a) every shareholder who desires to sell or transfer any share or shares shall give notice in writing to the board of such desire;\n(b) such notice shall be irrevocable and shall be deemed to appoint the board the proposing transferor's agent to sell such shares at a price to be agreed upon or determined by an independent expert at a fair price.\n(4) Offer to Shareholders and Consequent Sale: The board shall immediately give notice to each of the remaining shareholders offering the shares pro-rata.\n(5) Shares on Offer not Taken up by Shareholders: In the event of all such shares not being sold within 60 days, the party desiring to sell shall be at liberty to sell the shares to persons who are not shareholders at no less than the offered price."),
        ("4. Refusal to register transfers", "Directors' Right to refuse registration: Subject to compliance with the provisions of section 81 of the Act, the board may refuse or delay the registration of any transfer of any share to any person whether an existing shareholder or not –\n(a) if so required by law;\n(b) if the board acting in good faith decides in its sole discretion that registration of the Transfer would not be in the best interests of the company and/or any of its shareholders."),
        ("5. Purchase or other acquisition of own shares", "(1) Authority to Acquire Own Shares: for the purposes of section 65 of the Act, the company is expressly authorised to purchase or otherwise acquire shares issued by it."),
        ("6. Calls on shares and forfeiture of shares", "Calls on shares and forfeiture of shares shall be conducted in accordance with the Seventh Schedule to the Act."),
        ("7. Intellectual Property & Confidentiality (Special Provision)", "(1) All IP Assets developed by or for the company are and shall remain the exclusive property of the company.\n(2) Every director, employee, and contractor shall execute an IP Assignment and Confidentiality Agreement.\n(3) Where any shareholder also acts as legal counsel or advisor, they must execute a written Conflict of Interest Waiver before rendering advice."),
        ("8. Shareholders meetings", "Shareholders meetings shall be conducted in accordance with Part A of this Constitution."),
        ("9. Directors", "(1) The directors of the company shall be such person or persons as may be appointed from time to time by ordinary resolution.\n(2) The Founder, Kennedy Jr. Zibo Thebe, shall be a permanent director for so long as he holds shares in the company, and may not be removed by ordinary resolution."),
        ("10. Remuneration of directors", "The remuneration of directors shall be determined in accordance with section 157 of the Act."),
        ("11. Proceedings of directors meetings", "Proceedings of meetings of the Board of Directors shall be conducted in accordance with Part B of this Constitution."),
        ("12. Managing Director", "The directors may appoint one or more amongst themselves to the office of managing director (Chief Executive Officer)."),
        ("13. Dividends", "(1) All dividends shall be authorised by the board pursuant to section 58 of the Act.\n(2) Special Restriction: No dividends shall be declared or paid while any institutional grant disbursement agreement contains a restriction on profit distributions, or while any subsidiary has outstanding shareholder loans owed to the company."),
        ("14. Winding up", "(1) Upon the winding up of the company, the surplus assets shall be distributed among the shareholders in proportion to their shareholding.\n(2) Special Provision: The Founder shall have a right of first offer to acquire the IP Assets at fair market value before those assets are offered to any third party upon winding up."),
        ("15. One person companies and companies in which all shareholders are directors", "If at any time the company, for a continuous period exceeding six months is a one-person company, or is a company in which all shareholders also hold office as director, then, for so long as such circumstance continues, new shares may be issued by unanimous resolution and separate meetings of shareholders and directors need not be held.")
    ]

    for title, body in sections:
        p = doc.add_paragraph()
        p.add_run(title + "\n").bold = True
        p.add_run(body)

    doc.add_page_break()

    # Part A
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("PART A\nPROCEEDINGS AT MEETINGS OF SHAREHOLDERS\n").bold = True

    part_a = [
        ("16. Chairperson", "If the directors have elected a chairperson of the board, he or she shall chair the meeting of shareholders."),
        ("17. Notice of meetings", "Written notice of the time and place of a meeting of shareholders shall be sent to every shareholder not less than 10 working days before the meeting."),
        ("18. Methods of holding meetings", "A meeting of shareholders may be held either by a number of shareholders being assembled together, or by means of audio/visual communication."),
        ("19. Quorum", "A quorum for a meeting of shareholders is present if shareholders or their proxies are present who are between them able to exercise a majority of the votes."),
        ("20. Voting", "(1) Every shareholder present in person or by proxy shall have one vote per share.\n(2) Special Voting Rights: The Founder's shares carry an additional weighted voting right such that, in respect of any matter that could reduce the Founder's control of the company (including issue of new shares, amendments to the constitution, or permanent IP transfers), the Founder's vote counts as three (3) votes per share held."),
        ("21. Proxies", "A shareholder may exercise the right to vote either by being present in person or by proxy using Annexure I of this constitution."),
        ("22. Postal votes", "A shareholder may exercise the right to vote at a meeting by casting a postal vote."),
        ("23. Minutes", "The board shall ensure that minutes are kept of all proceedings at meetings of shareholders."),
        ("24. Shareholder proposals", "A shareholder may give written notice to the board of a matter the shareholder proposes to raise for discussion or resolution."),
        ("25. Corporations may act by representative", "A body corporate which is a shareholder may appoint a representative to attend a meeting on its behalf.")
    ]

    for title, body in part_a:
        p = doc.add_paragraph()
        p.add_run(title + "\n").bold = True
        p.add_run(body)

    doc.add_page_break()

    # Part B
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("PART B\nPROCEEDINGS AT MEETINGS OF THE BOARD OF A COMPANY\n").bold = True

    part_b = [
        ("26. Chairperson", "The directors may elect one of their number as chairperson of the board."),
        ("27. Notice of Meeting", "Not less than two days' notice of a meeting of the board must be sent to every director."),
        ("28. Methods of holding meetings", "A meeting of the board may be held by directors assembled together or by audio/visual communication."),
        ("29. Quorum", "A quorum for a meeting of the board shall be fixed by the board and if not so fixed shall be by majority of the directors."),
        ("30. Voting", "Every director has one vote. A resolution of the board is passed if it is agreed to by all directors present without dissent or if a majority of the votes cast on it are in favour of it."),
        ("31. Minutes", "The board must ensure that minutes are kept of all proceedings at meetings of the board."),
        ("32. Resolution in writing", "A resolution in writing, signed or assented to by all directors then entitled to receive notice of a board meeting, is as valid and effective as if it had been passed at a meeting."),
        ("33. The Secretary", "The Secretary shall keep all meeting minutes, records of the company and shall maintain an up-to-date Register of Shareholders.")
    ]

    for title, body in part_b:
        p = doc.add_paragraph()
        p.add_run(title + "\n").bold = True
        p.add_run(body)

    # Part C
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("\nPART C\nSECTIONS OF THIS ACT THAT CONFER POWERS ON DIRECTORS THAT CANNOT BE DELEGATED\n").bold = True

    part_c = [
        ("34. Powers that cannot be delegated", "1. The powers exercised in the company shall be as contemplated in this Constitution, and such powers shall regulate and bind the company in terms of section 41 (a) and (b) of the Act.\n2. The following sections confer powers on directors which cannot be delegated as per section 129 of the Act:\n(a) section 50 (issue of shares);\n(b) sections 53 and 54 (consideration for the issue of shares);\n(c) section 58 (distributions);\n(d) section 61 (issue of shares in lieu of dividends);\n(e) section 66 (offers to acquire shares);\n(f) section 76 (provision of financial assistance);\n(g) section 184 (change of registered office);"),
        ("35. Controllers of the Company", "2.1 The powers exercised in the company shall be under the ultimate effective control of the Board, each Director and each Shareholder, and ultimately the beneficial owners who exercise ultimate effective control.\n2.2 Annexure II contains a list of natural persons having ultimate effective control over the company in terms of section 41(b) of the Act.")
    ]

    for title, body in part_c:
        p = doc.add_paragraph()
        p.add_run(title + "\n").bold = True
        p.add_run(body)

    doc.add_page_break()

    # Annexures
    p = doc.add_paragraph()
    p.add_run("Annexure I - Proxy Form\n").bold = True
    p.add_run("\nThe instrument appointing a proxy shall be in the following form:\n\nI/we ________________________________________ of ________________________________________, Botswana being shareholders of the above-named company hereby appoint ________________________________________ Or failing him/her ________________________________________ of ________________________________________. As my/our proxy to vote for me/us at the meeting of the company to be held on ____ / ____ / 202__ and at any adjournment of the meeting.\n\nSigned this ____ day of ________________ 202__.\n\nSignature: ___________________________________")

    doc.add_page_break()

    p = doc.add_paragraph()
    p.add_run("Annexure II - Controllers of the Company\n\nCONTROLLER'S FORM\n(Section 21(2)(c))\n\n").bold = True
    
    p.add_run("Name of Company: Kgosi Sovereign Holdings Proprietary Limited\n")
    p.add_run("Company Number: [Pending]\n\n")
    p.add_run("Important Note: If there is more than one controller, each of the controllers should fill in a separate form.\n\n")
    p.add_run("CONTROLLER'S DETAILS\n\n").bold = True
    
    p.add_run("Controller's Name:\n")
    p.add_run("Kennedy Jr. Zibo Thebe\n\n")
    p.add_run("Residential Address:\n")
    p.add_run("_________________________________________________________________\n_________________________________________________________________\n\n")
    p.add_run("Position in the Company / Nature of Association:\n")
    p.add_run("Founder, Director, and Chief Executive Officer\n\n")
    p.add_run("Percentage of Contribution Held:\n")
    p.add_run("100% of all issued shares at incorporation. A further 35,000 shares are authorised but unissued and reserved for future allotment.\n\n")
    p.add_run("Signature: ___________________________________\n\n")
    p.add_run("Date: ____ / ____ / 202__\n\n")
    
    p.add_run("IMPORTANT INFORMATION\n").bold = True
    p.add_run("* provide full names and residential address of every beneficial owner including amount to be paid or other consideration.\n* where the beneficial owner is a representative, managerial position must be disclosed.\n* where some shares are to be held by a foreign company, the identification of natural persons who own, hold shares and control the foreign company must be disclosed.\n* beneficial owner’s interest must be expressed in percentage.\n\n")
    p.add_run("Completed by: Kennedy Jr. Zibo Thebe\n")
    p.add_run("Postal Address: _________________________________________________")

    doc.save(r'C:\Users\R5 5600 GT\Desktop\Kgosi_Constitution_Official.docx')

if __name__ == "__main__":
    create_doc()
