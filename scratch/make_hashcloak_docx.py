import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import os

def main():
    md_path = r"C:\Users\R5 5600 GT\.gemini\antigravity\scratch\resource_command\02_Engagements\Hashcloak\Hashcloak_Call_Cheat_Sheet.md"
    docx_path = r"C:\Users\R5 5600 GT\.gemini\antigravity\scratch\resource_command\02_Engagements\Hashcloak\Hashcloak_Call_Cheat_Sheet.docx"
    
    if not os.path.exists(md_path):
        print(f"MD file not found: {md_path}")
        return
        
    doc = docx.Document()
    
    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    color_primary = RGBColor(31, 58, 96)   # Deep Blue
    color_secondary = RGBColor(96, 31, 31) # Deep Red
    color_charcoal = RGBColor(44, 44, 44)  # Body text
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line_stripped = line.strip()
        
        if line_stripped == "---":
            p = doc.add_paragraph()
            run = p.add_run("__________________________________________________________________")
            run.font.color.rgb = RGBColor(200, 200, 200)
            continue
            
        if line_stripped.startswith("# "):
            title = line_stripped.lstrip("# ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(20)
            run.font.color.rgb = color_primary
            continue
            
        if line_stripped.startswith("### "):
            title = line_stripped.lstrip("### ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            run = p.add_run(title)
            run.bold = True
            run.font.size = Pt(14)
            run.font.color.rgb = color_secondary
            continue
            
        if line_stripped:
            is_bullet = False
            if line_stripped.startswith("* ") or line_stripped.startswith("- "):
                is_bullet = True
            
            is_numbered = False
            num_match = re.match(r'^(\d+)\.\s(.*)', line_stripped)
            if num_match:
                is_numbered = True
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            
            if is_bullet:
                p.paragraph_format.left_indent = Inches(0.25)
                cleaned_line = "\u2022 " + line_stripped[2:]
            elif is_numbered:
                p.paragraph_format.left_indent = Inches(0.25)
                cleaned_line = num_match.group(1) + ". " + num_match.group(2)
            else:
                cleaned_line = line_stripped
                
            parts = re.split(r'(\*\*.*?\*\*)', cleaned_line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                else:
                    run = p.add_run(part)
                run.font.size = Pt(11)
                run.font.color.rgb = color_charcoal

    doc.save(docx_path)
    print(f"DOCX created successfully at: {docx_path}")

if __name__ == "__main__":
    main()
