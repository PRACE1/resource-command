import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re
import os

def set_cell_background(cell, fill_hex):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets internal padding of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text):
    """Adds a stylish callout box for blockquotes."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(5.8)
    
    # Left border styling in XML
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="36" w:space="0" w:color="1F3A60"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    set_cell_background(cell, "F4F6F9")
    set_cell_margins(cell, top=150, bottom=150, left=200, right=200)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    # Text run formatting
    run = p.add_run(text.strip('"').strip())
    run.font.name = 'Calibri'
    run.font.size = Pt(11)
    run.font.italic = True
    run.font.color.rgb = RGBColor(44, 44, 44)
    
    doc.add_paragraph() # spacing after table

def main():
    md_path = r"C:\Users\R5 5600 GT\.gemini\antigravity\brain\c991fd44-7e0e-447c-bc12-807ef3894b12\geopolitical_positioning.md"
    docx_path = r"C:\Users\R5 5600 GT\.gemini\antigravity\scratch\resource_command\Resource_Command_Geopolitical_Positioning.docx"
    
    if not os.path.exists(md_path):
        print(f"MD file not found: {md_path}")
        return
        
    doc = docx.Document()
    
    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styling variables
    color_primary = RGBColor(31, 58, 96)   # Deep Blue
    color_secondary = RGBColor(96, 31, 31) # Deep Red
    color_charcoal = RGBColor(44, 44, 44)  # Body text
    
    # Read Markdown
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_mermaid = False
    in_table = False
    table_lines = []
    
    i = 0
    while i < len(lines):
        line = lines[i]
        line_stripped = line.strip()
        
        # Skip Mermaid diagrams (Word can't render them directly, so we skip the block)
        if line_stripped.startswith("```mermaid"):
            in_mermaid = True
            i += 1
            continue
        elif in_mermaid and line_stripped.startswith("```"):
            in_mermaid = False
            i += 1
            continue
        elif in_mermaid:
            i += 1
            continue
            
        # Parse horizontal rule
        if line_stripped == "---":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run("__________________________________________________________________")
            run.font.color.rgb = RGBColor(200, 200, 200)
            i += 1
            continue
            
        # Parse blockquotes (Narratives/Rebuttals)
        if line_stripped.startswith(">"):
            # Gather all consecutive blockquote lines
            bq_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                bq_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            text = " ".join(bq_lines)
            add_callout(doc, text)
            continue
            
        # Parse tables
        if line_stripped.startswith("|"):
            table_lines.append(line_stripped)
            i += 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            
            # Format and generate Table in Word
            rows_data = []
            for tl in table_lines:
                # Split cell contents and ignore empty start/end parts
                cells = [c.strip() for c in tl.split("|")[1:-1]]
                rows_data.append(cells)
                
            # Filter out separator row (e.g. | :--- | :--- |)
            rows_data = [r for r in rows_data if not all(re.match(r'^:?-+:?$', c) for c in r)]
            
            if rows_data:
                num_rows = len(rows_data)
                num_cols = len(rows_data[0])
                table = doc.add_table(rows=num_rows, cols=num_cols)
                table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
                table.autofit = True
                
                for r_idx, row_cells in enumerate(rows_data):
                    is_header = (r_idx == 0)
                    for c_idx, cell_val in enumerate(row_cells):
                        cell = table.cell(r_idx, c_idx)
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(6)
                        p.paragraph_format.space_after = Pt(6)
                        
                        # Strip bold syntax from text inside tables
                        cleaned_val = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_val)
                        run = p.add_run(cleaned_val)
                        run.font.name = 'Calibri'
                        
                        if is_header:
                            run.bold = True
                            run.font.color.rgb = RGBColor(255, 255, 255)
                            set_cell_background(cell, "1F3A60") # Header color
                        else:
                            run.font.color.rgb = color_charcoal
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F4F6F9") # Zebra striping
                                
                        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                
                doc.add_paragraph() # Add margin after table
            table_lines = []
            continue
            
        # Parse H1
        if line_stripped.startswith("# "):
            title = line_stripped.lstrip("# ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(title)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.color.rgb = color_primary
            i += 1
            continue
            
        # Parse H2
        if line_stripped.startswith("## "):
            title = line_stripped.lstrip("## ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(16)
            run.font.color.rgb = color_primary
            i += 1
            continue
            
        # Parse H3
        if line_stripped.startswith("### "):
            title = line_stripped.lstrip("### ").strip()
            # Remove MD bold syntax from header
            title = title.replace("**", "")
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(title)
            run.bold = True
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.color.rgb = color_secondary
            i += 1
            continue
            
        # Parse bullet points and standard paragraphs
        if line_stripped:
            # Handle bullet list
            is_bullet = False
            bullet_char = ""
            if line_stripped.startswith("* ") or line_stripped.startswith("- "):
                is_bullet = True
                bullet_char = "* " if line_stripped.startswith("* ") else "- "
                
            # Handle numbered list
            is_numbered = False
            num_match = re.match(r'^(\d+)\.\s(.*)', line_stripped)
            if num_match:
                is_numbered = True
                
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            
            if is_bullet:
                p.paragraph_format.left_indent = Inches(0.25)
                cleaned_line = line_stripped[2:]
                # Bold parsing inside list
                parts = re.split(r'(\*\*.*?\*\*)', cleaned_line)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        run = p.add_run(part)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    run.font.color.rgb = color_charcoal
            elif is_numbered:
                p.paragraph_format.left_indent = Inches(0.25)
                num_prefix = num_match.group(1) + ". "
                cleaned_line = num_match.group(2)
                
                run_num = p.add_run(num_prefix)
                run_num.bold = True
                run_num.font.name = 'Calibri'
                run_num.font.size = Pt(11)
                run_num.font.color.rgb = color_primary
                
                # Bold parsing inside numbered text
                parts = re.split(r'(\*\*.*?\*\*)', cleaned_line)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        run = p.add_run(part)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    run.font.color.rgb = color_charcoal
            else:
                # Standard paragraph
                # Bold parsing
                parts = re.split(r'(\*\*.*?\*\*)', line_stripped)
                for part in parts:
                    if part.startswith("**") and part.endswith("**"):
                        run = p.add_run(part[2:-2])
                        run.bold = True
                    else:
                        run = p.add_run(part)
                    run.font.name = 'Calibri'
                    run.font.size = Pt(11)
                    run.font.color.rgb = color_charcoal
                    
        i += 1
        
    doc.save(docx_path)
    print(f"DOCX created successfully at: {docx_path}")

if __name__ == "__main__":
    main()
